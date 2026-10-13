"""Importación de una hoja .xlsx como fotografía del stock por ubicación."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
import unicodedata
from uuid import uuid4

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook

from .inventory_services import apply_stock
from .models import Branch, Category, Product, StockBalance, StockMovement, StorageLocation

HEADERS = [
    "codigo", "nombre", "categoria", "marca", "unidad",
    "precio", "costo", "sucursal", "ubicacion", "cantidad", "minimo",
]
ALIASES = {"sku": "codigo", "producto": "nombre", "existencias": "cantidad",
           "stock": "cantidad", "almacen": "sucursal", "repisa": "ubicacion"}


def excel_template():
    book = Workbook()
    sheet = book.active
    sheet.title = "Inventario"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="193959")
    widths = [19, 36, 22, 18, 14, 13, 13, 25, 24, 14, 12]
    for index, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _normalize(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(char for char in text if not unicodedata.combining(char)).replace(" ", "_")


def _number(value, label, row, *, integer=False):
    if value in ("", None):
        return 0 if integer else Decimal("0")
    try:
        d = Decimal(str(value))
        if not d.is_finite() or d < 0 or (integer and d != d.to_integral_value()):
            raise InvalidOperation
        if integer:
            if d > 100000000:
                raise InvalidOperation
            return int(d)
        if d > Decimal("9999999999.99"):
            raise InvalidOperation
        return d.quantize(Decimal("0.01"))
    except (ValueError, TypeError, InvalidOperation):
        raise ValidationError(f"Fila {row}: {label} debe ser un número válido y no negativo.")


def _parse_excel(upload):
    if not upload.name.lower().endswith(".xlsx"):
        raise ValidationError("Sube un archivo .xlsx (Excel moderno).")
    if upload.size > 2 * 1024 * 1024:
        raise ValidationError("El archivo no debe superar 2 MB.")
    try:
        book = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError("No se pudo abrir el Excel. Revisa que sea un .xlsx válido.") from exc
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        raw_head = next(rows, None)
        if not raw_head:
            raise ValidationError("El Excel no contiene encabezados.")
        names = [_normalize(h) for h in raw_head]
        names = [ALIASES.get(name, name) for name in names]
        required = {"codigo", "nombre", "sucursal", "cantidad"}
        if not required.issubset(set(names)):
            raise ValidationError(
                "Faltan columnas obligatorias: codigo, nombre, sucursal, cantidad."
            )
        columns = {name: idx for idx, name in enumerate(names)}
        parsed = []
        for rownum, raw in enumerate(rows, 2):
            if all(value in ("", None) for value in raw):
                continue
            if len(parsed) >= 500:
                raise ValidationError("Importa como máximo 500 filas por archivo.")
            def cell(field):
                position = columns.get(field)
                return raw[position] if position is not None and position < len(raw) else None
            sku = str(cell("codigo") or "").strip().upper()
            name = str(cell("nombre") or "").strip()
            branch = str(cell("sucursal") or "").strip()
            location = str(cell("ubicacion") or "").strip() or "General"
            if not sku or not name or not branch:
                raise ValidationError(f"Fila {rownum}: codigo, nombre y sucursal son obligatorios.")
            if any(len(v) > limit for v, limit in (
                (sku, 64), (name, 180), (branch, 120), (location, 70)
            )):
                raise ValidationError(f"Fila {rownum}: algún texto supera el límite permitido.")
            parsed.append({
                "sku": sku, "name": name, "branch": branch, "location": location,
                "category": str(cell("categoria") or "").strip()[:90],
                "brand": str(cell("marca") or "").strip()[:100],
                "unit": str(cell("unidad") or "").strip()[:30] or "pieza",
                "quantity": _number(cell("cantidad"), "cantidad", rownum, integer=True),
                "min_stock": _number(cell("minimo"), "minimo", rownum, integer=True) if cell("minimo") not in ("", None) else 5,
                "price": _number(cell("precio"), "precio", rownum),
                "cost": _number(cell("costo"), "costo", rownum),
            })
        if not parsed:
            raise ValidationError("No se encontraron productos para importar.")
        return parsed
    finally:
        book.close()


@transaction.atomic
def import_excel(*, upload, actor):
    """Cantidad representa el stock FINAL de la ubicación, no una compra extra."""
    data = _parse_excel(upload)
    created = 0
    movements = 0
    for item in data:
        warehouse = Branch.objects.filter(code="FERMO-ALMACEN", is_active=True).first()
        if warehouse is not None:
            if item["branch"].casefold() not in {warehouse.name.casefold(), warehouse.code.casefold()}:
                raise ValidationError(
                    "Excel: utiliza únicamente Almacén FERMO en la columna sucursal. "
                    "No se permiten nuevas tiendas."
                )
            branch = warehouse
        else:
            # Conserva compatibilidad con pruebas e instalaciones sin inicializar.
            branch = Branch.objects.filter(name__iexact=item["branch"]).first()
            if branch is None:
                branch = Branch.objects.create(
                    code=f"IMP-{uuid4().hex[:12].upper()}", name=item["branch"]
                )
        location, _ = StorageLocation.objects.get_or_create(
            branch=branch, code=item["location"]
        )
        category = None
        if item["category"]:
            category, _ = Category.objects.get_or_create(name=item["category"])
        product, is_new = Product.objects.get_or_create(
            sku=item["sku"],
            defaults={
                "name": item["name"], "category": category, "brand": item["brand"],
                "unit": item["unit"], "price": item["price"], "cost": item["cost"],
                "min_stock": item["min_stock"],
            }
        )
        created += int(is_new)
        balance = StockBalance.objects.select_for_update().filter(
            product=product, location=location
        ).first()
        previous = balance.quantity if balance else 0
        difference = item["quantity"] - previous
        if difference:
            apply_stock(
                product=product, location=location, delta=difference,
                actor=actor, kind=StockMovement.IMPORT,
                note="Sincronización de existencia desde Excel",
                reference="XLSX-" + upload.name[:65],
            )
            movements += 1
    return len(data), created, movements
