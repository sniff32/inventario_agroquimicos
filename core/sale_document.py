"""Emite órdenes de venta Word (.docx), con logo corporativo y precios congelados."""
from io import BytesIO
from pathlib import Path
import re
from xml.etree import ElementTree as ET

from django.conf import settings
from django.utils import timezone
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

COMPANY_NAME = "Tienda de Agroquímicos"
LOGO = Path(settings.BASE_DIR) / "static" / "img" / "agro-logo.svg"


def _logo_png():
    """Renderiza la traza SVG del símbolo gráfico configurado con Pillow, sin binarios externos."""
    if not LOGO.exists():
        return None
    from PIL import Image, ImageChops, ImageDraw
    root = ET.parse(LOGO).getroot()
    path = next((e for e in root.iter() if e.tag.endswith("path")), None)
    if path is None:
        return None
    d = path.get("d", "")
    width, height, scale = 800, 300, 2
    mask = Image.new("1", (width * scale, height * scale), 0)
    for contour in re.findall(r"M([^M]+?)Z", d):
        coords = re.findall(r"(\d+),(\d+)", contour)
        if len(coords) < 3:
            continue
        poly = Image.new("1", mask.size, 0)
        ImageDraw.Draw(poly).polygon(
            [(int(x) * scale, int(y) * scale) for x, y in coords], fill=1
        )
        mask = ImageChops.logical_xor(mask, poly)
    image = Image.new("RGB", mask.size, "#ffffff")
    image.paste((38, 117, 72), (0, 0, mask.width, mask.height), mask)
    out = BytesIO()
    image.save(out, format="PNG", optimize=True)
    out.seek(0)
    return out


def _money(value):
    return f"$ {value:,.2f}"


def _shade(cell, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shade)


def _run(paragraph, text, *, size=9, bold=False, color=(38, 49, 60)):
    part = paragraph.add_run(str(text))
    part.bold = bold
    part.font.size = Pt(size)
    part.font.color.rgb = RGBColor(*color)
    return part


def build_sale_docx(sale):
    """Documento comercial interno: importes guardados en la venta, sin recalcular precios."""
    red = (191, 30, 44)
    charcoal = (40, 44, 49)
    muted = (100, 108, 116)
    cream = "FAF3F0"
    dark_red = "9F1E2C"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.57)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(*charcoal)
    normal.paragraph_format.space_after = Pt(4)

    logo_stream = _logo_png()
    brand = document.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if logo_stream:
        brand.add_run().add_picture(logo_stream, width=Inches(3.10))
    else:
        _run(brand, COMPANY_NAME, size=15, bold=True, color=red)

    line = document.add_table(rows=1, cols=1)
    line.rows[0].height = Inches(0.055)
    line.cell(0, 0).text = ""
    _shade(line.cell(0, 0), dark_red)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(12)
    _run(heading, "ORDEN DE VENTA", size=17, bold=True, color=charcoal)
    document.add_paragraph().paragraph_format.space_after = Pt(1)

    created = timezone.localtime(sale.created_at)
    fields = [
        ("FOLIO", sale.number, "EMISIÓN", created.strftime("%d/%m/%Y · %H:%M")),
        ("CLIENTE", sale.customer_name, "TIPO", sale.get_sale_type_display()),
        ("SUCURSAL", sale.branch.name, "SALIDA / ENTREGA", sale.delivery_date.strftime("%d/%m/%Y")),
        ("CONTACTO", sale.customer_contact or "No proporcionado", "FORMA DE PAGO", sale.get_payment_term_display()),
        ("VENCIMIENTO", sale.due_date.strftime("%d/%m/%Y") if sale.due_date else "—", "MONEDA", "MXN"),
    ]
    info = document.add_table(rows=5, cols=2)
    info.autofit = False
    for row, values in zip(info.rows, fields):
        for idx, (label, value) in enumerate(((values[0], values[1]), (values[2], values[3]))):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(8)
            _run(p, label + "\n", size=7, bold=True, color=red)
            _run(p, value, size=9, bold=True, color=charcoal)

    document.add_paragraph().paragraph_format.space_after = Pt(0)
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = False
    headers = ("SKU", "PRODUCTO · UBICACIÓN", "UDS.", "P. BASE", "DESC.", "P. FINAL", "IMPORTE")
    widths = [0.72, 2.30, 0.47, 0.84, 0.54, 0.87, 0.97]
    for col, (cell, title) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[col])
        _shade(cell, dark_red)
        cell.text = ""
        p = cell.paragraphs[0]
        _run(p, title, size=7.2, bold=True, color=(255, 255, 255))
    for idx, item in enumerate(sale.items.select_related("location").all()):
        cells = table.add_row().cells
        values = (
            item.sku_snapshot,
            item.name_snapshot + " (" + item.get_sale_type_display() + ")\n" + item.location.code,
            str(item.quantity),
            _money(item.base_unit_price),
            f"{item.discount_percent}%",
            _money(item.unit_price),
            _money(item.line_total),
        )
        for col, (cell, val) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[col])
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2 == 1:
                _shade(cell, cream)
            p = cell.paragraphs[0]
            if col > 1:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(p, val, size=7.6, color=charcoal)

    document.add_paragraph().paragraph_format.space_after = Pt(0)
    total_table = document.add_table(rows=3, cols=2)
    total_table.autofit = False
    for row, label, value, accent in [
        (total_table.rows[0], "Subtotal antes de descuentos", _money(sale.gross_subtotal), False),
        (total_table.rows[1], "Descuento por volumen", "− " + _money(sale.discount_total), False),
        (total_table.rows[2], "TOTAL FINAL · MXN", _money(sale.subtotal), True),
    ]:
        left, right = row.cells
        if accent:
            _shade(left, dark_red)
            _shade(right, dark_red)
        for cell, value, align in ((left, label, WD_ALIGN_PARAGRAPH.LEFT),
                                  (right, value, WD_ALIGN_PARAGRAPH.RIGHT)):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = align
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)
            _run(p, value, size=10 if accent else 8.5, bold=accent,
                 color=(255, 255, 255) if accent else charcoal)

    if sale.note:
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        _run(p, "OBSERVACIONES  ", bold=True, color=red)
        _run(p, sale.note, color=charcoal)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    _run(p, "REGISTRADO POR  ", size=7.5, bold=True, color=red)
    _run(p, sale.created_by.get_full_name() or sale.created_by.username, size=8.5)
    p = document.add_paragraph()
    _run(p,
         "AGROQUÍMICOS · Documento de control comercial interno, no es factura ni CFDI. "
         "Los importes no incluyen cálculo o desglose de impuestos; "
         "validar condiciones comerciales y fiscales antes de su uso oficial.",
         size=7, color=muted)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer, COMPANY_NAME + " · " + sale.number,
         size=7, bold=True, color=red)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
